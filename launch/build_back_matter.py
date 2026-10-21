#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
The doors in the books: pages to paste into the back of The Second Number, Book One and into
The Third Option, plus the Third Option fixes before it ships. Writes launch/Back_Matter_Pages.docx.

  python3 launch/build_back_matter.py
(The QR codes come from game/print_kit/build_where_next.py, which writes launch/qr/.)
"""
import os
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Back_Matter_Pages.docx")
QR = os.path.join(HERE, "qr")
GOLD = RGBColor(0x8A, 0x5A, 0x00)
GREY = RGBColor(0x6B, 0x66, 0x76)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Georgia"
st.font.size = Pt(11.5)


def note(text):
    """an instruction to the author, never printed in the book"""
    p = doc.add_paragraph()
    r = p.add_run("▸ " + text)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY


def para(text, bold=False, italic=False, size=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def door(title, text, url, qr=None):
    p = doc.add_paragraph()
    p.add_run(title).bold = True
    p.add_run(" " + text)
    u = doc.add_paragraph()
    r = u.add_run(url)
    r.bold = True
    r.font.color.rgb = GOLD
    if qr:
        doc.add_picture(os.path.join(QR, qr), width=Inches(1.1))


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- cover sheet
doc.add_heading("Back-matter pages for both books", 0)
para("Copy each page into the book's Word file where its note says. The grey notes in italics are for you only: "
     "don't paste them. The QR codes are images: copy them in too, or use the files in launch/qr/.")
para("Every address here is on getcivicarmor.com. They work once the hub is live, and they keep working "
     "even if where you sell the books changes: you change hub/config.js, never the printed page.", italic=True)
page_break()

# ---------------------------------------------------------------- the novel
doc.add_heading("THE SECOND NUMBER, BOOK ONE", 1)
note("Page 1 of 2. It goes after “A note on the numbers” and the reading list, on its own page. "
     "Gloria names Civic Armor in the last chapter; this page tells the reader it's real.")
doc.add_heading("Civic Armor is real", 2)
para("Gloria named it on a porch on Delmar: armor, ’cause it covers you. Civic, ’cause it ain’t yours.")
para("What Marcus could see, you can learn. Not with a Sight. With a method, and with the people on your block.")
door("The Third Option.", "The Civic Armor guide for your own job: how to recognize, document and prove what’s happening at work, "
     "before you ever see a lawyer.", "getcivicarmor.com/guide", "qr_third_option.png")
door("The Second Number, the game.", "A cooperative game for 1 to 5 players. Read the deals with your family, your class or your church. "
     "Free to play in English and Spanish.", "getcivicarmor.com/play", "qr_play.png")
door("Book Two, and everything after.", "New books, the boxed game and the Civic Armor course: join the list.",
     "getcivicarmor.com")
page_break()

note("Page 2 of 2: About the author. A draft from what's in The Third Option's introduction; make it yours.")
doc.add_heading("About the author", 2)
para("Lamont Ragland II spent more than three years studying real discrimination cases, one at a time, and studying "
     "history the same way: how the smaller side, the one with less power, still finds a way to win. What he found "
     "became Civic Armor, a method for reading the machine in your own life, and The Second Number, a novel and a "
     "cooperative board game that teach the same thing at the kitchen table.")
para("He is not a lawyer. He is the person who helps you understand what happened before you ever walk into a "
     "lawyer’s office.")
para("getcivicarmor.com", bold=True, color=GOLD)
page_break()

# ---------------------------------------------------------------- The Third Option
doc.add_heading("THE THIRD OPTION", 1)
note("Fix 1: the copyright page. It says “Published by Force Multiplier LLC” and “an instrument of Force Multiplier LLC.” "
     "Use the exact legal name on your state filing (the funnel says Force Multiplier HQ LLC) in both places. "
     "Replace those lines with:")
para("A Civic Armor Guide · Published by Force Multiplier HQ LLC, doing business as Civic Armor")
para("Copyright © 2026 by Lamont Ragland II. All rights reserved. Civic Armor™ is a trademark of Force Multiplier HQ LLC.")
para("getcivicarmor.com")
page_break()

note("Fix 2: the bracketed “PENDING ATTORNEY VERIFICATION … remove before publication” note near the end must not ship. "
     "Best: have an employment attorney write those topics. If the book ships first, delete the bracketed note and put "
     "this honest box in its place:")
doc.add_heading("What this edition doesn’t cover", 2)
para("Some rules depend on your state, the size of your employer and your exact situation, and this edition doesn’t cover "
     "them in detail. Ask a licensed employment attorney about each of these, and ask early, because some of the "
     "deadlines are short:")
for item in ["how long you have to file a lawsuit after the EEOC sends a notice of right to sue (often 90 days)",
             "severance agreements, and the review time older workers get before signing one in a layoff",
             "whether your employer is big enough to be covered by a particular law",
             "your state’s rules on recording conversations",
             "your duty to look for other work after losing a job",
             "every protected class, including national origin, color, genetic information, sexual orientation and gender identity",
             "the particular shape of retaliation claims and failure-to-accommodate claims"]:
    doc.add_paragraph(item, style="List Bullet")
para("This book helps you see and prepare. A lawyer confirms and acts.", italic=True)
page_break()

note("Fix 3: “The Next Step” page. Keep its text and add this at the bottom, so the reader has somewhere to go. "
     "getcivicarmor.com/course sends people to the course section of the hub (a waitlist until the course opens).")
door("The Civic Armor course:", "", "getcivicarmor.com/course", "qr_course.png")
page_break()

note("New last page, after “The Next Step.” It's for readers who'd like the story: Renée's readers meet Marcus.")
doc.add_heading("If you want the story", 2)
para("The method in this book has a world behind it.")
door("The Second Number, Book One.", "A novel. At eighteen, Marcus Ellison signed a plea so he could go home. "
     "Then he wakes up on graduation day, able to see the hidden cost in every deal on his block, and has to decide "
     "what to do with it without becoming the thing he’s reading.", "getcivicarmor.com/book", "qr_book.png")
door("The Second Number, the game.", "Read the deals with your family, your class or your church. Free to play.",
     "getcivicarmor.com/play")

doc.save(OUT)
print(f"wrote {OUT}")
